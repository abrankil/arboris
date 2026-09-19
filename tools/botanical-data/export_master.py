from __future__ import annotations

import hashlib
import json
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
MASTER = (
    ROOT
    / "data"
    / "source"
    / "Base_botanica_Pokedex_flora_Master_2.0_FINAL.xlsx"
)
OUTPUT_DIR = ROOT / "data" / "botanical"

EXPORT_MAP_SHEET = "Exportar_JSON"
SCHEMA_SHEET = "Diccionario_Campos"
METADATA_SHEET = "Metadatos"


class ExportError(RuntimeError):
    pass


def clean_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def is_yes(value):
    if isinstance(value, bool):
        return value
    text = clean_text(value)
    if text is None:
        return False
    return text.casefold() in {"si", "sí", "yes", "true", "1"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rows_as_dicts(ws):
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return []

    headers = [clean_text(value) for value in header_row]
    records = []

    for excel_row_number, row in enumerate(rows, start=2):
        if all(value is None or clean_text(value) is None for value in row):
            continue

        record = {
            header: value
            for header, value in zip(headers, row)
            if header is not None
        }
        record["__excel_row__"] = excel_row_number
        records.append(record)

    return records


def load_export_plan(workbook):
    plan_rows = rows_as_dicts(workbook[EXPORT_MAP_SHEET])
    plan = []

    for row in plan_rows:
        if not is_yes(row.get("exportar")):
            continue

        sheet = clean_text(row.get("hoja"))
        json_key = clean_text(row.get("json_key"))

        if not sheet or not json_key:
            raise ExportError(
                f"{EXPORT_MAP_SHEET} fila {row['__excel_row__']}: "
                "una exportación activa requiere hoja y json_key."
            )

        plan.append(
            {
                "order": row.get("orden"),
                "sheet": sheet,
                "json_key": json_key,
                "filename": f"{json_key}.json",
                "transformation": clean_text(row.get("transformacion")),
            }
        )

    if not plan:
        raise ExportError("Exportar_JSON no contiene exportaciones activas.")

    return sorted(
        plan,
        key=lambda item: (
            item["order"] if isinstance(item["order"], (int, float)) else 10**9,
            item["sheet"],
        ),
    )


def load_schema(workbook):
    schema_rows = rows_as_dicts(workbook[SCHEMA_SHEET])
    by_sheet = defaultdict(list)

    for row in schema_rows:
        sheet = clean_text(row.get("hoja"))
        field = clean_text(row.get("campo"))
        json_type = clean_text(row.get("json_tipo"))
        nullable = clean_text(row.get("nullable"))

        if not sheet or not field or not json_type:
            raise ExportError(
                f"{SCHEMA_SHEET} fila {row['__excel_row__']}: "
                "hoja, campo y json_tipo son obligatorios."
            )

        by_sheet[sheet].append(
            {
                "field": field,
                "json_type": json_type,
                "nullable": is_yes(nullable),
                "row": row["__excel_row__"],
            }
        )

    return dict(by_sheet)


def normalize_boolean(value, *, sheet, field, row_number):
    if isinstance(value, bool):
        return value

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value in {0, 1}:
            return bool(value)

    text = clean_text(value)
    if text is not None:
        normalized = text.casefold()
        if normalized in {"true", "1", "si", "sí", "yes"}:
            return True
        if normalized in {"false", "0", "no"}:
            return False

    raise ExportError(
        f"{sheet} fila {row_number}, campo {field}: "
        f"valor booleano inválido {value!r}."
    )


def normalize_date(value, *, sheet, field, row_number):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()

    text = clean_text(value)
    if text is None:
        return None

    try:
        return datetime.fromisoformat(text).date().isoformat()
    except ValueError as exc:
        raise ExportError(
            f"{sheet} fila {row_number}, campo {field}: "
            f"fecha inválida {value!r}."
        ) from exc


def normalize_number(value, *, sheet, field, row_number):
    if isinstance(value, bool):
        raise ExportError(
            f"{sheet} fila {row_number}, campo {field}: "
            "un booleano no es un número válido."
        )

    if isinstance(value, (int, float)):
        return value

    text = clean_text(value)
    if text is None:
        return None

    try:
        number = float(text)
    except ValueError as exc:
        raise ExportError(
            f"{sheet} fila {row_number}, campo {field}: "
            f"número inválido {value!r}."
        ) from exc

    return int(number) if number.is_integer() else number


def transform_value(
    value,
    *,
    json_type,
    nullable,
    separator,
    sheet,
    field,
    row_number,
):
    if value is None or clean_text(value) is None:
        if not nullable and json_type == "array<string>":
            return []
        return None

    if json_type == "array<string>":
        if isinstance(value, (list, tuple)):
            return [str(item).strip() for item in value if clean_text(item) is not None]
        return [
            item.strip()
            for item in str(value).split(separator)
            if item.strip()
        ]

    if json_type == "date":
        return normalize_date(
            value,
            sheet=sheet,
            field=field,
            row_number=row_number,
        )

    if json_type == "boolean":
        return normalize_boolean(
            value,
            sheet=sheet,
            field=field,
            row_number=row_number,
        )

    if json_type == "number":
        return normalize_number(
            value,
            sheet=sheet,
            field=field,
            row_number=row_number,
        )

    if json_type == "string|number":
        if isinstance(value, bool):
            return str(value)
        if isinstance(value, (int, float)):
            return value
        return clean_text(value)

    if json_type == "string":
        return clean_text(value)

    raise ExportError(
        f"{SCHEMA_SHEET}: tipo JSON no soportado {json_type!r} "
        f"para {sheet}.{field}."
    )


def metadata_to_object(workbook):
    rows = rows_as_dicts(workbook[METADATA_SHEET])
    metadata = {}

    for row in rows:
        key = clean_text(row.get("clave"))
        if not key:
            continue
        metadata[key] = row.get("valor")

    return metadata


def sheet_to_records(workbook, sheet_name, field_specs, separator):
    ws = workbook[sheet_name]
    rows = ws.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration:
        return []

    headers = [clean_text(value) for value in header_row]
    header_to_index = {
        header: index
        for index, header in enumerate(headers)
        if header is not None
    }

    required_fields = [spec["field"] for spec in field_specs]
    missing_columns = [
        field
        for field in required_fields
        if field not in header_to_index
    ]

    if missing_columns:
        raise ExportError(
            f"{sheet_name}: faltan columnas declaradas en {SCHEMA_SHEET}: "
            + ", ".join(missing_columns)
        )

    records = []

    for excel_row_number, row in enumerate(rows, start=2):
        if all(value is None or clean_text(value) is None for value in row):
            continue

        record = {}

        for spec in field_specs:
            field = spec["field"]
            raw_value = row[header_to_index[field]]
            value = transform_value(
                raw_value,
                json_type=spec["json_type"],
                nullable=spec["nullable"],
                separator=separator,
                sheet=sheet_name,
                field=field,
                row_number=excel_row_number,
            )

            if value is None and not spec["nullable"]:
                raise ExportError(
                    f"{sheet_name} fila {excel_row_number}, campo {field}: "
                    "el esquema declara nullable=No pero el valor está vacío."
                )

            record[field] = value

        records.append(record)

    return records


def json_text(data):
    return json.dumps(
        data,
        ensure_ascii=False,
        indent=2,
    ) + "\n"


def write_json(filename, data):
    path = OUTPUT_DIR / filename
    path.write_text(json_text(data), encoding="utf-8")
    print(f"Created: {path.relative_to(ROOT)}")


def main():
    if not MASTER.exists():
        raise ExportError(f"Master no encontrado: {MASTER}")

    print(f"Reading: {MASTER.relative_to(ROOT)}")

    workbook = load_workbook(
        MASTER,
        data_only=True,
        read_only=True,
    )

    try:
        required_control_sheets = {
            METADATA_SHEET,
            EXPORT_MAP_SHEET,
            SCHEMA_SHEET,
        }
        missing_control = required_control_sheets - set(workbook.sheetnames)
        if missing_control:
            raise ExportError(
                "Faltan hojas de control: " + ", ".join(sorted(missing_control))
            )

        export_plan = load_export_plan(workbook)
        schema = load_schema(workbook)
        metadata = metadata_to_object(workbook)

        separator = str(metadata.get("multi_state_separator") or "|")

        for item in export_plan:
            if item["sheet"] not in workbook.sheetnames:
                raise ExportError(
                    f"{EXPORT_MAP_SHEET}: hoja exportable inexistente {item['sheet']!r}."
                )

        datasets = {}

        for item in export_plan:
            sheet_name = item["sheet"]
            json_key = item["json_key"]

            if sheet_name == METADATA_SHEET:
                continue

            if sheet_name not in schema:
                raise ExportError(
                    f"{sheet_name}: no tiene campos definidos en {SCHEMA_SHEET}."
                )

            datasets[json_key] = sheet_to_records(
                workbook,
                sheet_name,
                schema[sheet_name],
                separator,
            )

        metadata_output = dict(metadata)
        metadata_output["source_file"] = MASTER.name
        metadata_output["source_sha256"] = sha256_file(MASTER)

        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

        for item in export_plan:
            if item["sheet"] == METADATA_SHEET:
                continue
            write_json(
                item["filename"],
                datasets[item["json_key"]],
            )

        metadata_filename = next(
            item["filename"]
            for item in export_plan
            if item["sheet"] == METADATA_SHEET
        )
        write_json(metadata_filename, metadata_output)

    finally:
        workbook.close()

    print("\nMaster 2.0 export completed successfully.")


if __name__ == "__main__":
    try:
        main()
    except ExportError as exc:
        print(f"\nEXPORT FAILED\n- {exc}", file=sys.stderr)
        raise SystemExit(1)
