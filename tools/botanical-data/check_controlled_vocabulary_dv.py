from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from controlled_vocabulary import (
    ControlledVocabularyError,
    load_spec,
    resolve_field,
)


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_MASTER = (
    ROOT
    / "data"
    / "source"
    / "Base_botanica_Pokedex_flora_Master_2.0_FINAL.xlsx"
)

RECORD_KEYS = {
    "Especie_Caracter": ("species_id", "caracter_id"),
    "Ecologia_Especie": ("ecology_fact_id",),
}

FORBIDDEN_LIST_DV_FIELDS = {
    "Especie_Caracter.notas",
}


class DVCheckError(RuntimeError):
    pass


def header_map(ws) -> dict[str, int]:
    result: dict[str, int] = {}

    for cell in ws[1]:
        if cell.value is None:
            continue

        name = str(cell.value)

        if name in result:
            raise DVCheckError(
                f"{ws.title}: duplicate header {name!r}."
            )

        result[name] = cell.column

    return result


def record_rows(ws, keys: tuple[str, ...]) -> list[int]:
    headers = header_map(ws)

    missing = [key for key in keys if key not in headers]
    if missing:
        raise DVCheckError(
            f"{ws.title}: missing record-key headers {missing!r}."
        )

    rows: list[int] = []

    for row_number in range(2, ws.max_row + 1):
        if any(
            ws.cell(row_number, headers[key]).value is not None
            for key in keys
        ):
            rows.append(row_number)

    return rows


def dv_applies_to_cell(dv, row: int, column: int) -> bool:
    for cell_range in dv.ranges.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(
            str(cell_range)
        )

        if (
            min_row <= row <= max_row
            and min_col <= column <= max_col
        ):
            return True

    return False


def resolve_list_tokens(workbook, ws, dv) -> tuple[str, ...]:
    formula = dv.formula1

    if not isinstance(formula, str) or not formula:
        raise DVCheckError(
            f"{ws.title}: list DV has no resolvable formula1."
        )

    if formula.startswith('"') and formula.endswith('"'):
        inner = formula[1:-1]

        if inner == "":
            return tuple()

        return tuple(inner.split(","))

    reference = formula[1:] if formula.startswith("=") else formula

    if "!" in reference:
        sheet_part, range_part = reference.rsplit("!", 1)
        sheet_name = sheet_part.strip("'").replace("''", "'")
    else:
        sheet_name = ws.title
        range_part = reference

    if sheet_name not in workbook.sheetnames:
        raise DVCheckError(
            f"{ws.title}: unresolved DV sheet {sheet_name!r}."
        )

    try:
        min_col, min_row, max_col, max_row = range_boundaries(
            range_part.replace("$", "")
        )
    except ValueError as exc:
        raise DVCheckError(
            f"{ws.title}: unsupported or unresolved DV source "
            f"{formula!r}."
        ) from exc

    source_ws = workbook[sheet_name]
    values: list[str] = []

    for row in source_ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            if cell.value is None:
                raise DVCheckError(
                    f"{ws.title}: blank value in DV source "
                    f"{formula!r}."
                )

            if not isinstance(cell.value, str):
                raise DVCheckError(
                    f"{ws.title}: non-string value in DV source "
                    f"{formula!r}."
                )

            values.append(cell.value)

    return tuple(values)


def check_workbook(
    workbook,
    spec: dict[str, Any],
) -> list[str]:
    errors: list[str] = []

    for scoped_name in spec["scope"]:
        sheet_name, separator, field_name = scoped_name.partition(".")

        if not separator or not sheet_name or not field_name:
            raise DVCheckError(
                f"Invalid governed field reference {scoped_name!r}."
            )

        if sheet_name not in workbook.sheetnames:
            errors.append(
                f"{scoped_name}: missing sheet {sheet_name!r}."
            )
            continue

        ws = workbook[sheet_name]
        headers = header_map(ws)

        if field_name not in headers:
            errors.append(
                f"{scoped_name}: missing exact header {field_name!r}."
            )
            continue

        if sheet_name not in RECORD_KEYS:
            raise DVCheckError(
                f"No record-key contract for sheet {sheet_name!r}."
            )

        rows = record_rows(ws, RECORD_KEYS[sheet_name])
        column = headers[field_name]

        try:
            contract = resolve_field(scoped_name, spec)
        except ControlledVocabularyError as exc:
            raise DVCheckError(
                f"Cannot resolve governed field {scoped_name!r}: {exc}"
            ) from exc

        for row_number in rows:
            applicable = [
                dv
                for dv in ws.data_validations.dataValidation
                if dv_applies_to_cell(dv, row_number, column)
            ]

            location = (
                f"{sheet_name}!"
                f"{get_column_letter(column)}{row_number}"
            )

            if len(applicable) != 1:
                errors.append(
                    f"{location}: expected exactly one DV, "
                    f"found {len(applicable)}."
                )
                continue

            dv = applicable[0]

            if dv.type != "list":
                errors.append(
                    f"{location}: DV type {dv.type!r}, "
                    f"expected 'list'."
                )
                continue

            try:
                tokens = resolve_list_tokens(workbook, ws, dv)
            except DVCheckError as exc:
                errors.append(f"{location}: {exc}")
                continue

            expected_tokens = tuple(contract["allowed_values"])

            if tokens != expected_tokens:
                errors.append(
                    f"{location}: DV tokens {tokens!r}, "
                    f"expected {expected_tokens!r}."
                )

            if bool(dv.allow_blank) != contract["nullable"]:
                errors.append(
                    f"{location}: allow_blank={bool(dv.allow_blank)!r}, "
                    f"expected nullable={contract['nullable']!r}."
                )

    for scoped_name in sorted(FORBIDDEN_LIST_DV_FIELDS):
        sheet_name, _, field_name = scoped_name.partition(".")

        if sheet_name not in workbook.sheetnames:
            continue

        ws = workbook[sheet_name]
        headers = header_map(ws)

        if field_name not in headers:
            continue

        if sheet_name not in RECORD_KEYS:
            raise DVCheckError(
                f"No record-key contract for sheet {sheet_name!r}."
            )

        rows = record_rows(ws, RECORD_KEYS[sheet_name])
        column = headers[field_name]

        for row_number in rows:
            applicable_lists = [
                dv
                for dv in ws.data_validations.dataValidation
                if (
                    dv.type == "list"
                    and dv_applies_to_cell(
                        dv,
                        row_number,
                        column,
                    )
                )
            ]

            if applicable_lists:
                location = (
                    f"{sheet_name}!"
                    f"{get_column_letter(column)}{row_number}"
                )
                errors.append(
                    f"{location}: forbidden list DV on "
                    f"{scoped_name}."
                )

    return errors


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "workbook",
        nargs="?",
        type=Path,
        default=DEFAULT_MASTER,
    )
    parser.add_argument(
        "--spec",
        type=Path,
        default=None,
    )
    args = parser.parse_args()

    try:
        spec = load_spec(args.spec)
    except ControlledVocabularyError as exc:
        raise SystemExit(
            f"CONTROLLED VOCABULARY SPEC INVALID: {exc}"
        ) from exc

    workbook = load_workbook(
        args.workbook,
        data_only=False,
        read_only=False,
    )

    try:
        errors = check_workbook(workbook, spec)
    except DVCheckError as exc:
        raise SystemExit(f"DV CHECKER ERROR: {exc}") from exc
    finally:
        workbook.close()

    print("?RBORIS ? CONTROLLED VOCABULARY DV CHECK")
    print("=" * 60)
    print(f"Workbook: {args.workbook}")
    print(f"Spec: {args.spec or 'default'}")
    print(f"Errors: {len(errors)}")

    for error in errors:
        print(f"ERROR: {error}")

    if errors:
        raise SystemExit(1)

    print("OK ? controlled-vocabulary DV contract satisfied.")


if __name__ == "__main__":
    main()
