from __future__ import annotations

import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from controlled_vocabulary import (
    ControlledVocabularyError,
    load_spec,
    resolve_field,
    validate_spec,
)


ROOT = Path(__file__).resolve().parents[2]
MASTER = (
    ROOT
    / "data"
    / "source"
    / "Base_botanica_Pokedex_flora_Master_2.0_FINAL.xlsx"
)
BOTANICAL_DIR = ROOT / "data" / "botanical"

EXPORT_MAP_SHEET = "Exportar_JSON"
SCHEMA_SHEET = "Diccionario_Campos"
METADATA_SHEET = "Metadatos"

FORBIDDEN_BOTANICAL_STATES = {
    "no_se",
    "no_sé",
    "no_observable",
    "no_aplica",
    "unknown",
    "desconocido",
}

ALLOWED_CHARACTER_STATUS = {
    "activo",
    "retirado",
    "pendiente_revision",
}

DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


class ValidationError(RuntimeError):
    pass


def clean_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def is_yes(value):
    if isinstance(value, bool):
        return value
    text = clean_text(value)
    if text is None:
        return False
    return text.casefold() in {"si", "sí", "yes", "true", "1"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rows_as_dicts(ws):
    rows = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return []

    headers = [clean_text(value) for value in header_row]
    records = []

    for excel_row_number, row in enumerate(rows, start=2):
        if all(value is None or clean_text(value) is None for value in row):
            continue

        record = {
            header: value
            for header, value in zip(headers, row)
            if header is not None
        }
        record["__excel_row__"] = excel_row_number
        records.append(record)

    return records


def load_export_plan(workbook):
    plan = []

    for row in rows_as_dicts(workbook[EXPORT_MAP_SHEET]):
        if not is_yes(row.get("exportar")):
            continue

        sheet = clean_text(row.get("hoja"))
        json_key = clean_text(row.get("json_key"))

        if not sheet or not json_key:
            raise ValidationError(
                f"{EXPORT_MAP_SHEET} fila {row['__excel_row__']}: "
                "exportación activa sin hoja/json_key."
            )

        plan.append(
            {
                "order": row.get("orden"),
                "sheet": sheet,
                "json_key": json_key,
                "filename": f"{json_key}.json",
                "transformation": clean_text(row.get("transformacion")),
            }
        )

    return sorted(
        plan,
        key=lambda item: (
            item["order"] if isinstance(item["order"], (int, float)) else 10**9,
            item["sheet"],
        ),
    )


def load_schema(workbook):
    by_sheet = defaultdict(list)

    for row in rows_as_dicts(workbook[SCHEMA_SHEET]):
        sheet = clean_text(row.get("hoja"))
        field = clean_text(row.get("campo"))
        json_type = clean_text(row.get("json_tipo"))

        if not sheet and not field and not json_type:
            continue

        if not sheet or not field or not json_type:
            raise ValidationError(
                f"{SCHEMA_SHEET} fila {row['__excel_row__']}: "
                "hoja, campo y json_tipo son obligatorios."
            )

        by_sheet[sheet].append(
            {
                "field": field,
                "json_type": json_type,
                "nullable": is_yes(row.get("nullable")),
                "relation": clean_text(row.get("relacion")),
                "row": row["__excel_row__"],
            }
        )

    return dict(by_sheet)


def metadata_from_master(workbook):
    metadata = {}

    for row in rows_as_dicts(workbook[METADATA_SHEET]):
        key = clean_text(row.get("clave"))
        if key:
            metadata[key] = row.get("valor")

    return metadata


def normalize_date(value, *, sheet, field, row_number):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()

    text = clean_text(value)
    if text is None:
        return None

    try:
        return datetime.fromisoformat(text).date().isoformat()
    except ValueError as exc:
        raise ValidationError(
            f"{sheet} fila {row_number}, campo {field}: fecha inválida {value!r}."
        ) from exc


def normalize_boolean(value, *, sheet, field, row_number):
    if isinstance(value, bool):
        return value

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value in {0, 1}:
            return bool(value)

    text = clean_text(value)
    if text is not None:
        normalized = text.casefold()
        if normalized in {"true", "1", "si", "sí", "yes"}:
            return True
        if normalized in {"false", "0", "no"}:
            return False

    raise ValidationError(
        f"{sheet} fila {row_number}, campo {field}: booleano inválido {value!r}."
    )


def normalize_number(value, *, sheet, field, row_number):
    if isinstance(value, bool):
        raise ValidationError(
            f"{sheet} fila {row_number}, campo {field}: booleano no válido como número."
        )

    if isinstance(value, (int, float)):
        return value

    text = clean_text(value)
    if text is None:
        return None

    try:
        number = float(text)
    except ValueError as exc:
        raise ValidationError(
            f"{sheet} fila {row_number}, campo {field}: número inválido {value!r}."
        ) from exc

    return int(number) if number.is_integer() else number


def transform_value(value, *, spec, separator, sheet, row_number):
    json_type = spec["json_type"]
    nullable = spec["nullable"]
    field = spec["field"]

    if value is None or clean_text(value) is None:
        if not nullable and json_type == "array<string>":
            return []
        return None

    if json_type == "array<string>":
        if isinstance(value, (list, tuple)):
            return [
                str(item).strip()
                for item in value
                if clean_text(item) is not None
            ]
        return [
            item.strip()
            for item in str(value).split(separator)
            if item.strip()
        ]

    if json_type == "date":
        return normalize_date(
            value,
            sheet=sheet,
            field=field,
            row_number=row_number,
        )

    if json_type == "boolean":
        return normalize_boolean(
            value,
            sheet=sheet,
            field=field,
            row_number=row_number,
        )

    if json_type == "number":
        return normalize_number(
            value,
            sheet=sheet,
            field=field,
            row_number=row_number,
        )

    if json_type == "string|number":
        if isinstance(value, bool):
            return str(value)
        if isinstance(value, (int, float)):
            return value
        return clean_text(value)

    if json_type == "string":
        return clean_text(value)

    raise ValidationError(
        f"{SCHEMA_SHEET}: tipo no soportado {json_type!r} para {sheet}.{field}."
    )


def expected_records(workbook, sheet_name, field_specs, separator):
    ws = workbook[sheet_name]
    rows = ws.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration:
        return []

    headers = [clean_text(value) for value in header_row]
    header_to_index = {
        header: index
        for index, header in enumerate(headers)
        if header is not None
    }

    missing = [
        spec["field"]
        for spec in field_specs
        if spec["field"] not in header_to_index
    ]
    if missing:
        raise ValidationError(
            f"{sheet_name}: faltan columnas declaradas en {SCHEMA_SHEET}: "
            + ", ".join(missing)
        )

    records = []

    for excel_row_number, row in enumerate(rows, start=2):
        if all(value is None or clean_text(value) is None for value in row):
            continue

        record = {}

        for spec in field_specs:
            field = spec["field"]
            value = transform_value(
                row[header_to_index[field]],
                spec=spec,
                separator=separator,
                sheet=sheet_name,
                row_number=excel_row_number,
            )

            if value is None and not spec["nullable"]:
                raise ValidationError(
                    f"{sheet_name} fila {excel_row_number}, campo {field}: "
                    "nullable=No pero el valor está vacío."
                )

            record[field] = value

        records.append(record)

    return records


def load_json(path: Path):
    try:
        with path.open("r", encoding="utf-8") as file:
            return json.load(file)
    except FileNotFoundError as exc:
        raise ValidationError(f"Falta archivo requerido: {path}") from exc
    except json.JSONDecodeError as exc:
        raise ValidationError(f"JSON inválido en {path}: {exc}") from exc


def find_duplicates(values):
    counts = Counter(values)
    return [value for value, count in counts.items() if count > 1]


def normalize_state(value):
    if not isinstance(value, str):
        return value
    return value.strip().casefold()


def validate_json_value(value, spec, *, location, errors):
    json_type = spec["json_type"]
    nullable = spec["nullable"]

    if value is None:
        if not nullable:
            errors.append(f"{location}: null no permitido.")
        return

    valid = True

    if json_type == "string":
        valid = isinstance(value, str)
    elif json_type == "number":
        valid = isinstance(value, (int, float)) and not isinstance(value, bool)
    elif json_type == "string|number":
        valid = (
            isinstance(value, str)
            or (isinstance(value, (int, float)) and not isinstance(value, bool))
        )
    elif json_type == "boolean":
        valid = isinstance(value, bool)
    elif json_type == "date":
        valid = isinstance(value, str) and bool(DATE_RE.fullmatch(value))
        if valid:
            try:
                date.fromisoformat(value)
            except ValueError:
                valid = False
    elif json_type == "array<string>":
        valid = isinstance(value, list) and all(isinstance(item, str) for item in value)
    else:
        errors.append(f"{location}: tipo de esquema no soportado {json_type!r}.")
        return

    if not valid:
        errors.append(
            f"{location}: esperaba {json_type}, recibió {type(value).__name__}."
        )


def validate_controlled_vocabulary_records(
    records,
    *,
    sheet_name,
    spec,
    errors,
):
    """Validate governed vocabulary fields without normalization or coercion."""

    try:
        validate_spec(spec)
    except ControlledVocabularyError as exc:
        raise ValidationError(
            f"Controlled-vocabulary specification invalid: {exc}"
        ) from exc

    governed_fields = []

    for scoped_name in spec["scope"]:
        prefix, separator, field_name = scoped_name.partition(".")

        if not separator or not prefix or not field_name:
            raise ValidationError(
                f"Invalid governed field reference: {scoped_name!r}"
            )

        if prefix == sheet_name:
            governed_fields.append((scoped_name, field_name))

    if not governed_fields:
        raise ValidationError(
            f"No controlled-vocabulary fields governed for sheet "
            f"{sheet_name!r}."
        )

    for record_number, record in enumerate(records, start=1):
        if not isinstance(record, dict):
            errors.append(
                f"{sheet_name} registro {record_number}: "
                f"esperaba objeto, recibi? {type(record).__name__}."
            )
            continue

        sid = record.get("species_id")
        cid = record.get("caracter_id")

        if sid is not None or cid is not None:
            location = f"{sid}/{cid}"
        else:
            location = f"{sheet_name} registro {record_number}"

        for scoped_name, field_name in governed_fields:
            try:
                contract = resolve_field(scoped_name, spec)
            except ControlledVocabularyError as exc:
                raise ValidationError(
                    f"Cannot resolve controlled vocabulary "
                    f"{scoped_name!r}: {exc}"
                ) from exc

            if field_name not in record:
                errors.append(
                    f"{location}: falta campo gobernado {field_name!r}."
                )
                continue

            value = record[field_name]

            if value is None:
                if not contract["nullable"]:
                    errors.append(
                        f"{location}.{field_name}: null no permitido."
                    )
                continue

            if not isinstance(value, str):
                errors.append(
                    f"{location}.{field_name}: esperaba string, "
                    f"recibi? {type(value).__name__}."
                )
                continue

            if value not in contract["allowed_values"]:
                errors.append(
                    f"{location}.{field_name}: valor no permitido "
                    f"{value!r}; permitidos="
                    f"{list(contract['allowed_values'])!r}."
                )


def first_difference(expected, actual):
    if len(expected) != len(actual):
        return f"cantidad de registros: Master={len(expected)}, JSON={len(actual)}"

    for index, (exp, got) in enumerate(zip(expected, actual), start=1):
        if exp != got:
            exp_keys = set(exp)
            got_keys = set(got) if isinstance(got, dict) else set()

            if not isinstance(got, dict):
                return f"registro {index}: esperaba objeto, recibió {type(got).__name__}"

            if exp_keys != got_keys:
                return (
                    f"registro {index}: campos distintos; "
                    f"faltan={sorted(exp_keys - got_keys)}, "
                    f"sobran={sorted(got_keys - exp_keys)}"
                )

            for key in exp:
                if exp[key] != got[key]:
                    return (
                        f"registro {index}, campo {key}: "
                        f"Master={exp[key]!r}, JSON={got[key]!r}"
                    )

    return None


def main():
    print("\nÁRBORIS — VALIDACIÓN MASTER 2.0 EXPORTADO")
    print("=" * 72)

    errors = []
    warnings = []

    if not MASTER.exists():
        print(f"\nERROR — Master no encontrado: {MASTER}")
        raise SystemExit(1)

    workbook = load_workbook(MASTER, data_only=True, read_only=True)

    try:
        required_control_sheets = {
            METADATA_SHEET,
            EXPORT_MAP_SHEET,
            SCHEMA_SHEET,
        }
        missing_control = required_control_sheets - set(workbook.sheetnames)
        if missing_control:
            raise ValidationError(
                "Faltan hojas de control: " + ", ".join(sorted(missing_control))
            )

        export_plan = load_export_plan(workbook)
        schema = load_schema(workbook)
        master_metadata = metadata_from_master(workbook)

        try:
            controlled_vocabulary_spec = load_spec()
        except ControlledVocabularyError as exc:
            raise ValidationError(
                f"Controlled-vocabulary specification invalid: {exc}"
            ) from exc

        controlled_vocabulary_bindings = {
            "controlled_vocabulary_spec_id":
                controlled_vocabulary_spec["specification_id"],
            "controlled_vocabulary_spec_version":
                controlled_vocabulary_spec["specification_version"],
        }

        for binding_key, expected_value in (
            controlled_vocabulary_bindings.items()
        ):
            actual_value = master_metadata.get(binding_key)

            if actual_value != expected_value:
                errors.append(
                    f"{METADATA_SHEET}.{binding_key}: "
                    f"Master={actual_value!r}, "
                    f"SPEC={expected_value!r}."
                )

        if errors:
            raise ValidationError(
                "controlled-vocabulary binding mismatch"
            )

        separator = str(master_metadata.get("multi_state_separator") or "|")

        expected_filenames = {item["filename"] for item in export_plan}
        expected_filenames.add("metadata.json")

        for filename in sorted(expected_filenames):
            if not (BOTANICAL_DIR / filename).exists():
                errors.append(f"Falta exportación requerida: {filename}")

        if errors:
            raise ValidationError("faltan archivos exportados")

        datasets = {}
        sheet_for_key = {}

        for item in export_plan:
            sheet_name = item["sheet"]
            json_key = item["json_key"]
            filename = item["filename"]

            if sheet_name == METADATA_SHEET:
                continue

            if sheet_name not in workbook.sheetnames:
                errors.append(
                    f"{EXPORT_MAP_SHEET}: hoja inexistente {sheet_name!r}."
                )
                continue

            if sheet_name not in schema:
                errors.append(
                    f"{sheet_name}: no tiene contrato en {SCHEMA_SHEET}."
                )
                continue

            actual = load_json(BOTANICAL_DIR / filename)
            datasets[json_key] = actual
            sheet_for_key[json_key] = sheet_name

            if not isinstance(actual, list):
                errors.append(f"{filename} debe contener un array JSON.")
                continue

            specs = schema[sheet_name]
            expected_fields = [spec["field"] for spec in specs]
            expected_field_set = set(expected_fields)

            for record_number, record in enumerate(actual, start=1):
                if not isinstance(record, dict):
                    errors.append(
                        f"{filename} registro {record_number}: debe ser un objeto."
                    )
                    continue

                actual_fields = set(record)
                missing_fields = expected_field_set - actual_fields
                extra_fields = actual_fields - expected_field_set

                if missing_fields:
                    errors.append(
                        f"{filename} registro {record_number}: "
                        f"faltan campos {sorted(missing_fields)}."
                    )
                if extra_fields:
                    errors.append(
                        f"{filename} registro {record_number}: "
                        f"sobran campos {sorted(extra_fields)}."
                    )

                for spec in specs:
                    if spec["field"] in record:
                        validate_json_value(
                            record[spec["field"]],
                            spec,
                            location=(
                                f"{filename} registro {record_number}."
                                f"{spec['field']}"
                            ),
                            errors=errors,
                        )

            expected = expected_records(
                workbook,
                sheet_name,
                specs,
                separator,
            )
            difference = first_difference(expected, actual)
            if difference:
                errors.append(
                    f"{filename} no reproduce exactamente {sheet_name}: {difference}."
                )

        exported_metadata = load_json(BOTANICAL_DIR / "metadata.json")
        if not isinstance(exported_metadata, dict):
            errors.append("metadata.json debe contener un objeto JSON.")
            exported_metadata = {}

        for key, expected_value in master_metadata.items():
            actual_value = exported_metadata.get(key)
            if actual_value != expected_value:
                errors.append(
                    f"metadata.json.{key}: Master={expected_value!r}, "
                    f"JSON={actual_value!r}."
                )

        expected_source_file = MASTER.name
        expected_sha = sha256_file(MASTER)

        if exported_metadata.get("source_file") != expected_source_file:
            errors.append(
                "metadata.json.source_file no corresponde al Master actual."
            )

        if exported_metadata.get("source_sha256") != expected_sha:
            errors.append(
                "metadata.json.source_sha256 no coincide con el SHA-256 del Master."
            )

        allowed_metadata_keys = set(master_metadata) | {
            "source_file",
            "source_sha256",
        }
        extra_metadata = set(exported_metadata) - allowed_metadata_keys
        if extra_metadata:
            warnings.append(
                "metadata.json contiene claves no declaradas por el contrato: "
                + ", ".join(sorted(extra_metadata))
            )

        species = datasets.get("species", [])
        characters = datasets.get("characters", [])
        species_characters = datasets.get("species_characters", [])
        sources = datasets.get("sources", [])

        validate_controlled_vocabulary_records(
            species_characters,
            sheet_name="Especie_Caracter",
            spec=controlled_vocabulary_spec,
            errors=errors,
        )
        glossary = datasets.get("glossary", [])
        photos = datasets.get("photos", [])
        model_errors = datasets.get("model_errors", [])

        # -----------------------------------------------------
        # Identificadores y claves primarias
        # -----------------------------------------------------
        species_ids = [row.get("species_id") for row in species if isinstance(row, dict)]
        character_ids = [row.get("caracter_id") for row in characters if isinstance(row, dict)]
        source_ids = [row.get("fuente_id") for row in sources if isinstance(row, dict)]
        glossary_ids = [row.get("termino_id") for row in glossary if isinstance(row, dict)]
        photo_ids = [row.get("photo_id") for row in photos if isinstance(row, dict)]
        error_ids = [row.get("error_id") for row in model_errors if isinstance(row, dict)]

        for label, values in [
            ("species_id", species_ids),
            ("caracter_id", character_ids),
            ("fuente_id", source_ids),
            ("termino_id", glossary_ids),
            ("photo_id", photo_ids),
            ("error_id", error_ids),
        ]:
            for value in find_duplicates(v for v in values if v):
                errors.append(f"{label} duplicado: {value}")

        species_id_set = {value for value in species_ids if value}
        source_id_set = {value for value in source_ids if value}
        character_by_id = {
            row.get("caracter_id"): row
            for row in characters
            if isinstance(row, dict) and row.get("caracter_id")
        }

        # -----------------------------------------------------
        # Caracteres, estados y estado_piloto
        # -----------------------------------------------------
        status_counts = Counter()

        for character in characters:
            if not isinstance(character, dict):
                continue

            cid = character.get("caracter_id")
            status = character.get("estado_piloto")
            status_counts[status] += 1

            if status not in ALLOWED_CHARACTER_STATUS:
                errors.append(
                    f"{cid}: estado_piloto inválido {status!r}."
                )

            states = character.get("estados_permitidos")
            if not isinstance(states, list):
                continue

            if not states:
                errors.append(f"{cid}: estados_permitidos no puede estar vacío.")

            normalized = [normalize_state(state) for state in states]

            for state in find_duplicates(normalized):
                errors.append(f"{cid}: estado permitido duplicado {state!r}.")

            for state in normalized:
                if state in FORBIDDEN_BOTANICAL_STATES:
                    errors.append(
                        f"{cid}: {state!r} no debe ser un estado botánico."
                    )

        computable_status = exported_metadata.get("computable_status")
        excluded_raw = exported_metadata.get("excluded_status")
        excluded_status = {
            item.strip()
            for item in str(excluded_raw or "").split(separator)
            if item.strip()
        }

        if computable_status not in ALLOWED_CHARACTER_STATUS:
            errors.append(
                f"metadata computable_status desconocido: {computable_status!r}."
            )

        if excluded_status - ALLOWED_CHARACTER_STATUS:
            errors.append(
                "metadata excluded_status contiene estados desconocidos: "
                + ", ".join(sorted(excluded_status - ALLOWED_CHARACTER_STATUS))
            )

        if computable_status in excluded_status:
            errors.append(
                "metadata computable_status también aparece en excluded_status."
            )

        # -----------------------------------------------------
        # Relaciones especie-caracter
        # -----------------------------------------------------
        seen_pairs = set()

        for index, row in enumerate(species_characters, start=1):
            if not isinstance(row, dict):
                continue

            sid = row.get("species_id")
            cid = row.get("caracter_id")
            fid = row.get("fuente_id")
            expected_states = row.get("estado_esperado")

            if sid not in species_id_set:
                errors.append(f"{sid}/{cid}: species_id inexistente.")

            if cid not in character_by_id:
                errors.append(f"{sid}/{cid}: caracter_id inexistente.")

            pair = (sid, cid)
            if pair in seen_pairs:
                errors.append(f"Relación duplicada: {sid}/{cid}")
            seen_pairs.add(pair)

            if fid not in source_id_set:
                errors.append(f"{sid}/{cid}: fuente_id inexistente {fid!r}.")

            if not isinstance(expected_states, list):
                continue

            if not expected_states:
                errors.append(f"{sid}/{cid}: estado_esperado no puede estar vacío.")

            normalized_expected = [normalize_state(state) for state in expected_states]

            for state in find_duplicates(normalized_expected):
                errors.append(
                    f"{sid}/{cid}: estado_esperado duplicado {state!r}."
                )

            for state in normalized_expected:
                if state in FORBIDDEN_BOTANICAL_STATES:
                    errors.append(
                        f"{sid}/{cid}: {state!r} no debe usarse como estado botánico."
                    )

            character = character_by_id.get(cid)
            if character:
                allowed_states = {
                    normalize_state(state)
                    for state in character.get("estados_permitidos", [])
                }
                for state in normalized_expected:
                    if state not in allowed_states:
                        errors.append(
                            f"{sid}/{cid}: estado {state!r} fuera de estados_permitidos."
                        )

        # -----------------------------------------------------
        # Otras claves foráneas
        # -----------------------------------------------------
        for row in glossary:
            if isinstance(row, dict) and row.get("fuente_id") not in source_id_set:
                errors.append(
                    f"Glosario {row.get('termino_id')}: fuente_id inexistente "
                    f"{row.get('fuente_id')!r}."
                )

        for row in photos:
            if isinstance(row, dict) and row.get("species_id") not in species_id_set:
                errors.append(
                    f"Foto {row.get('photo_id')}: species_id inexistente "
                    f"{row.get('species_id')!r}."
                )

        for row in model_errors:
            if isinstance(row, dict) and row.get("species_id_real") not in species_id_set:
                errors.append(
                    f"Error {row.get('error_id')}: species_id_real inexistente "
                    f"{row.get('species_id_real')!r}."
                )

        # -----------------------------------------------------
        # Cobertura informativa
        # -----------------------------------------------------
        relations_by_species = Counter(
            row.get("species_id")
            for row in species_characters
            if isinstance(row, dict) and row.get("species_id")
        )
        photos_by_species = Counter(
            row.get("species_id")
            for row in photos
            if isinstance(row, dict) and row.get("species_id")
        )

        for sid in sorted(species_id_set):
            if relations_by_species[sid] == 0:
                errors.append(f"{sid}: no tiene caracteres registrados.")
            if photos_by_species[sid] == 0:
                warnings.append(f"{sid}: no tiene evidencia fotográfica.")

        # -----------------------------------------------------
        # Reporte
        # -----------------------------------------------------
        print()
        print(f"Especies:             {len(species)}")
        print(f"Caracteres:           {len(characters)}")
        print(f"  activos:            {status_counts.get('activo', 0)}")
        print(f"  retirados:          {status_counts.get('retirado', 0)}")
        print(f"  pendiente revisión: {status_counts.get('pendiente_revision', 0)}")
        print(f"Relaciones especie:   {len(species_characters)}")
        print(f"Fuentes:              {len(sources)}")
        print(f"Términos glosario:    {len(glossary)}")
        print(f"Fotos:                {len(photos)}")
        print(f"Errores de modelo:    {len(model_errors)}")
        print(f"SHA-256 Master:       {expected_sha}")

        print("\nCobertura por especie:")
        for sid in sorted(species_id_set):
            print(
                f"  {sid}: {relations_by_species[sid]} caracteres | "
                f"{photos_by_species[sid]} fotos"
            )

    except ValidationError as exc:
        errors.append(str(exc))
    finally:
        workbook.close()

    print("\n" + "=" * 72)

    if warnings:
        print(f"\nADVERTENCIAS: {len(warnings)}")
        for warning in warnings:
            print(f"- {warning}")
    else:
        print("\nAdvertencias: 0")

    if errors:
        print(f"\nERRORES: {len(errors)}")
        for error in errors:
            print(f"- {error}")
        print("\nMASTER 2.0 EXPORTADO: VALIDACIÓN FALLIDA")
        raise SystemExit(1)

    print("\nErrores: 0")
    print("\nOK — MASTER 2.0 EXPORTADO VÁLIDO")
    print("Los JSON reproducen exactamente el Master Botánico 2.0 y su contrato de esquema.")


if __name__ == "__main__":
    main()
